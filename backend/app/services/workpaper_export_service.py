"""The engagement's workpaper as an Excel file: what was reported, what was recalculated
and how, the evidence behind it, what the reviewer decided, and the open work.

Every number comes from the same recalculation the screens show; nothing is computed
here. Sheets: Summary, Final outputs, Roll-up, Categories (the emissions roll-up, when the
recalculation has one), Data points, Working, Evidence, Reviewed values, Requests,
Findings, Left out.
"""

from __future__ import annotations

import io
import re
from datetime import UTC, date, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

_HEADER_FILL = PatternFill("solid", fgColor="E8EAF4")
_STATUS = {
    "within_tolerance": "Held",
    "exceeds_tolerance": "Exception",
    "not_reported": "Could not check (reported value not entered)",
    "no_evidence": "Could not check (no reviewed evidence)",
    "not_recalculated": "Out of scope (not recalculated yet)",
    "not_in_scope": "Not in scope (set in Setup)",
    "incomplete": "Not concluded (the total is missing categories)",
    "needs_unit": "Could not check (reported unit not confirmed: t or t CO2e)",
}
_REASON = {
    "lower_bound": "Already above the reported value, and the missing categories can only add",
    "reported_zero": "Reported as 0",
    "partial_total": "Missing categories could still change the result",
    "unit_unconfirmed": "Set whether the company reported tonnes of gas or t CO2e",
    "not_built": "Category not built yet",
    "no_evidence": "No reviewed evidence",
    "awaiting_inputs": "Evidence waiting on reviewer inputs",
    "set_out_of_scope": "Set out of scope in Setup",
}
_MISSING = {
    "not_built": "not built yet",
    "no_evidence": "no evidence",
    "awaiting_inputs": "waiting on inputs",
    "inputs_open": "questions open",
    "vehicle_groups_missing": "no vehicle groups for CH4/N2O",
    "gas_split_unknown": "blend's gas split unknown",
}
_SOURCE_STATUS = {
    "calculated": "Calculated",
    "partial": "Calculated, questions open",
    "awaiting_inputs": "Waiting on inputs",
    "confirmed_absent": "Confirmed absent (0)",
    "no_evidence": "No evidence",
    "not_built": "Not built yet",
}
_NODE_STATUS = {"complete": "Complete", "partial": "Partial", "not_available": "Not available"}
_LEFT_OUT = {
    "outside_reporting_period": "Outside the reporting period",
    "withdrawn": "Withdrawn by the reviewer",
    "duplicate": "Marked as a duplicate",
    "superseded": "Superseded by a newer document",
    "montreal_protocol_gas": "A CFC or HCFC (Montreal Protocol gas), reported outside Scope 1",
}


def _sheet(book: Workbook, title: str, header: list[str], rows: list[list[Any]], widths: list[int]) -> None:
    sheet = book.create_sheet(title)
    sheet.append(header)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
        cell.fill = _HEADER_FILL
    for row in rows:
        sheet.append([_cell(v) for v in row])
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=isinstance(cell.value, str) and len(cell.value) > 40)
    sheet.freeze_panes = "A2"


def _cell(value: Any) -> Any:
    if value is None or isinstance(value, (int, float, str)):
        return value
    return str(value)


def _when(value: Any) -> str:
    return str(value or "")[:19].replace("T", " ")


_MONTH_ABBR = ("Jan", "Feb", "Mar", "Apr", "May", "Jun", "Jul", "Aug", "Sep", "Oct", "Nov", "Dec")
_ISO_DAY = re.compile(r"^(\d{4})-(\d{2})-(\d{2})")


def _day(value: Any) -> str:
    """One date as a person reads it, "Oct 8, 2026" (en-US); anything else as it was given."""
    if isinstance(value, date):
        return f"{_MONTH_ABBR[value.month - 1]} {value.day}, {value.year}"
    text = str(value)
    match = _ISO_DAY.match(text)
    if not match:
        return text
    year, month, day = (int(part) for part in match.groups())
    if not 1 <= month <= 12 or not 1 <= day <= 31:
        return text
    return f"{_MONTH_ABBR[month - 1]} {day}, {year}"


def format_reporting_period(start: Any, end: Any) -> str:
    """The reporting period in the words the Output screen uses (frontend formatReportingPeriod),
    so the screen and its download agree: "Oct 8, 2026 – Dec 31, 2026", or the missing date
    named in its place while the engagement is still being set up."""

    if start and end:
        return f"{_day(start)} – {_day(end)}"
    if start:
        return f"{_day(start)} – end date not set"
    if end:
        return f"start date not set – {_day(end)}"
    return "Reporting period not set"


def _missing(items: list[dict[str, Any]] | None, labels: dict[str, str]) -> str:
    return "; ".join(f"{labels.get(m.get('source', ''), m.get('source'))}: "
                     f"{_MISSING.get(str(m.get('reason')), m.get('reason'))}" for m in items or [])


def _rollup_sheets(book: Workbook, rollup: dict[str, Any]) -> None:
    """Final outputs, the Steps 1a/1b/2 nodes, and each category's state."""

    categories = rollup.get("categories") or {}
    labels = {key: str(c.get("label") or key) for key, c in categories.items()}
    _sheet(
        book, "Final outputs",
        ["Workbook row", "Data point", "Requirement", "Unit", "Reported", "Recalculated", "Difference",
         "Difference %", "Within threshold", "Outcome", "Why", "Completeness", "Missing", "Counted as 0"],
        [[o.get("workbook_row"), o.get("label"), o.get("requirement"), o.get("unit"), o.get("reported"),
          o.get("recalculated"), o.get("difference"), o.get("difference_pct"),
          {True: "yes", False: "no"}.get(o.get("within_threshold"), ""),
          _STATUS.get(str(o.get("status")), o.get("status")), _REASON.get(str(o.get("reason")), o.get("reason")),
          _NODE_STATUS.get(str(o.get("completeness")), o.get("completeness")), _missing(o.get("missing"), labels),
          _missing(o.get("counted_as_zero"), labels)]
         for o in rollup.get("final_outputs") or []],
        [12, 52, 12, 10, 14, 14, 14, 13, 10, 40, 44, 14, 60, 60],
    )
    nodes = rollup.get("nodes") or {}
    steps = {step.get("key"): step.get("label") for step in rollup.get("steps") or []}
    order = [key for step in rollup.get("steps") or [] for key in step.get("node_keys") or []]
    _sheet(
        book, "Roll-up",
        ["Step", "Line", "t CO2e", "t of gas", "Status", "Adds up", "Missing", "Notes", "Counted as 0"],
        [[steps.get(nodes[key].get("step"), nodes[key].get("step")), nodes[key].get("label"), nodes[key].get("co2e_t"),
          nodes[key].get("mass_t"), _NODE_STATUS.get(str(nodes[key].get("status")), nodes[key].get("status")),
          nodes[key].get("formula"), _missing(nodes[key].get("missing"), labels), ", ".join(nodes[key].get("flags") or []),
          _missing(nodes[key].get("counted_as_zero"), labels)]
         for key in order if key in nodes],
        [40, 52, 14, 12, 14, 52, 60, 40, 60],
    )
    _sheet(
        book, "Categories",
        ["Category", "Scope", "Workbook row", "Status", "Open questions", "Reviewer's reason it is absent", "Notes"],
        [[c.get("label"), c.get("scope"), c.get("workbook_row"), _SOURCE_STATUS.get(str(c.get("status")), c.get("status")),
          c.get("open_items"), c.get("absent_reason"), ", ".join(c.get("flags") or [])]
         for c in categories.values()],
        [48, 10, 12, 26, 14, 44, 40],
    )


def build_workbook(
    *,
    engagement: dict[str, Any],
    recalculation: dict[str, Any],
    register: dict[str, list[dict[str, Any]]],
    items: list[dict[str, Any]],
    exported_by: str,
) -> bytes:
    book = Workbook()
    book.remove(book.active)
    settings = engagement.get("settings") or {}
    totals = recalculation.get("scope_totals") or {}

    rollup = recalculation.get("rollup") or {}

    def total(name: str) -> float | None:
        part = totals.get(name)
        return part["total_co2e"] if part else None

    def node(name: str) -> Any:
        value = (rollup.get("nodes") or {}).get(name) or {}
        return value.get("co2e_t") if value.get("status") == "complete" else (
            f"{value.get('co2e_t'):.4f} (partial)" if value.get("co2e_t") is not None else "Not available")

    summary: list[list[Any]] = [
        ["Engagement", engagement.get("name")],
        ["Company", engagement.get("client_name")],
        ["Reporting period", format_reporting_period(engagement.get("reporting_period_start"),
                                                     engagement.get("reporting_period_end"))],
        ["Regulation", settings.get("regulation")],
        ["Assurance level", settings.get("assuranceLevel")],
        ["Assurance standard", settings.get("assuranceStandard")],
        ["Boundary approach", settings.get("boundaryApproach")],
        ["Materiality tolerance (%)", recalculation.get("tolerance_pct")],
        ["GWP basis", recalculation.get("gwp_set")],
        ["Emission factors", "EPA GHG Emission Factors Hub 2025"],
        ["Scope 1 recalculated (t CO2e)", total("scope1")],
        ["Scope 2 location-based recalculated (t CO2e)", total("scope2_location")],
        ["Scope 2 market-based recalculated (t CO2e)", total("scope2_market")],
        *([["Inventory total, location-based (t CO2e)", node("inventory_location_total")],
           ["Inventory total, market-based (t CO2e)", node("inventory_market_total")]] if rollup else []),
        ["Documents with reviewed values", recalculation.get("documents_considered")],
        ["Status", "Draft working paper, not an assurance opinion"],
        ["Exported", f"{datetime.now(UTC):%Y-%m-%d %H:%M} UTC by {exported_by}"],
    ]
    _sheet(book, "Summary", ["Item", "Value"], summary, [44, 70])
    if rollup:
        _rollup_sheets(book, rollup)

    _sheet(
        book, "Data points",
        ["Workbook row", "Data point", "Unit", "Reported", "Recalculated", "Difference", "Difference %",
         "Outcome", "Inputs incomplete"],
        [[c.get("workbook_row"), c.get("data_point"), c.get("unit"), c.get("reported"), c.get("recalculated"),
          c.get("difference"), c.get("difference_pct"), _STATUS.get(str(c.get("status")), c.get("status")),
          "yes" if c.get("inputs_incomplete") else ""]
         for c in recalculation.get("comparisons") or []],
        [12, 52, 10, 14, 14, 14, 13, 44, 16],
    )

    working: list[list[Any]] = []
    for key, category in (recalculation.get("categories") or {}).items():
        if not category:
            continue
        for line in category.get("lines") or []:
            for step in line.get("steps") or []:
                inputs = "; ".join(
                    f"{name} = {i.get('value')} {i.get('unit') or ''}".strip()
                    + (f" [{i.get('source')}]" if i.get("source") else "")
                    for name, i in (step.get("inputs") or {}).items()
                )
                working.append([key, line.get("key"), step.get("label"), step.get("formula"), inputs,
                                step.get("result"), step.get("result_unit")])
        for step in category.get("steps") or []:
            inputs = "; ".join(f"{name} = {i.get('value')} {i.get('unit') or ''}".strip()
                               for name, i in (step.get("inputs") or {}).items())
            working.append([key, "(category total)", step.get("label"), step.get("formula"), inputs,
                            step.get("result"), step.get("result_unit")])
    _sheet(book, "Working", ["Category", "Line", "Step", "Formula", "Inputs (with source)", "Result", "Unit"],
           working, [30, 26, 30, 30, 90, 16, 10])

    _sheet(
        book, "Evidence",
        ["File", "Document type", "Uploaded", "Uploaded by", "Values extracted"],
        [[d.get("file_name"), d.get("canonical_type_id") or "Not identified", _when(d.get("uploaded_at")),
          d.get("uploaded_by"), d.get("candidate_count")] for d in register.get("documents") or []],
        [48, 22, 20, 30, 16],
    )
    _sheet(
        book, "Reviewed values",
        ["File", "Record", "Field", "Extracted", "Unit", "Decision", "Reviewed value", "Reviewed unit",
         "Reviewer", "Reviewed at", "Page", "Source text"],
        [[v.get("file_name"), v.get("record_index"), v.get("field"), v.get("extracted"), v.get("unit"),
          v.get("decision"), v.get("reviewed_value"), v.get("reviewed_unit"), v.get("reviewer"),
          _when(v.get("reviewed_at")), v.get("page"), v.get("snippet")] for v in register.get("values") or []],
        [40, 8, 30, 18, 10, 14, 16, 12, 30, 20, 6, 60],
    )

    asks = [i["payload"] for i in items if i.get("kind") == "ask"]
    _sheet(
        book, "Requests",
        ["Item", "Type", "What is needed", "Facility", "Period", "State", "Raised by", "Raised at", "Request no."],
        [[(a.get("source") or {}).get("itemLabel"), a.get("type"), a.get("whatIsNeeded"),
          (a.get("source") or {}).get("facilityName"), (a.get("source") or {}).get("period"), a.get("state"),
          a.get("raisedByName"), _when(a.get("raisedAt")), a.get("requestNumber")] for a in asks],
        [40, 18, 60, 22, 22, 14, 18, 20, 12],
    )
    findings = [i["payload"] for i in items if i.get("kind") == "finding"]
    _sheet(
        book, "Findings",
        ["Data point", "Finding", "Cause", "Magnitude", "Status", "Registered by", "Registered at"],
        [[f.get("dataPointLabel"), f.get("sentence"), f.get("cause"), f.get("magnitude"), f.get("status"),
          f.get("registeredBy"), _when(f.get("registeredAt"))] for f in findings],
        [36, 60, 28, 16, 12, 18, 20],
    )
    _sheet(
        book, "Left out",
        ["Document / record", "Why it is not counted", "Dates"],
        [[e.get("record"), _LEFT_OUT.get(str(e.get("reason")), e.get("reason")), e.get("period")]
         for e in recalculation.get("excluded") or []],
        [56, 36, 26],
    )

    buffer = io.BytesIO()
    book.save(buffer)
    return buffer.getvalue()
