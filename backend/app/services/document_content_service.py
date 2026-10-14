"""The content of an uploaded spreadsheet, CSV or Word/text file, for the source pane.

A PDF or image is shown as a page picture with the value outlined on it. These files have
no page, so the pane draws their content instead: a grid for a workbook or CSV, lines for
Word and text. The numbering is the parsers' own, so an extracted value's source reference
points at the right spot:

- Excel: sheet name and cell ("D2"), as ``openpyxl`` reports them.
- CSV: the row number in the file, counting blank rows (the CSV parser's ``r<N>``).
- Word / text: the n-th non-empty line (the text parser's block ``b<N>``).

Large files are cut to a readable size and say so (``truncated``).
"""

from __future__ import annotations

import csv
from pathlib import Path
from typing import Any

from backend.app.adapters.parsers.docx_parser import DocxError, docx_text
from backend.app.adapters.parsers.excel_parser import _coerce_cell_value
from backend.app.services.parser_service import ParserService

MAX_SHEETS = 30
MAX_ROWS = 2000  # per sheet (rows with content)
MAX_COLUMNS = 60
MAX_LINES = 5000
MAX_CELL_CHARS = 500


class UnsupportedContent(ValueError):
    """The file is shown as pages (PDF, image), or cannot be shown at all."""


class UnreadableContent(ValueError):
    """The file could not be opened."""


def _cell(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value)
    if not text.strip():
        return None
    return text if len(text) <= MAX_CELL_CHARS else text[:MAX_CELL_CHARS] + "…"


def _workbook(path: Path) -> dict[str, Any]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except Exception as exc:  # noqa: BLE001 - any load failure is "unreadable"
        raise UnreadableContent("The workbook could not be opened.") from exc
    sheets: list[dict[str, Any]] = []
    truncated = False
    try:
        for index, worksheet in enumerate(workbook.worksheets):
            if index >= MAX_SHEETS:
                truncated = True
                break
            rows: list[dict[str, Any]] = []
            width = 0
            for row in worksheet.iter_rows():
                placed: dict[int, str] = {}
                number: int | None = None
                for cell in row:  # placed by the cell's own column, as its "D2" reference says
                    text = _cell(_coerce_cell_value(cell.value))
                    column = getattr(cell, "column", None)
                    if text is None or not isinstance(column, int):
                        continue
                    number = number or getattr(cell, "row", None)
                    if column > MAX_COLUMNS:
                        truncated = True
                        continue
                    placed[column - 1] = text
                if not placed:
                    continue
                if len(rows) >= MAX_ROWS:
                    truncated = True
                    break
                cells = [placed.get(i) for i in range(max(placed) + 1)]
                width = max(width, len(cells))
                rows.append({"number": number, "cells": cells})
            if rows:
                sheets.append({"name": worksheet.title, "columns": width, "rows": rows})
    finally:
        workbook.close()
    return {"kind": "sheets", "sheets": sheets, "truncated": truncated}


def _decoded(path: Path) -> str:
    try:
        return path.read_text(encoding="utf-8")
    except UnicodeDecodeError:
        return path.read_text(encoding="utf-8", errors="replace")


def _csv(path: Path) -> dict[str, Any]:
    rows: list[dict[str, Any]] = []
    width = 0
    truncated = False
    for number, values in enumerate(csv.reader(_decoded(path).splitlines()), start=1):
        if not any(str(v).strip() for v in values):
            continue
        if len(rows) >= MAX_ROWS:
            truncated = True
            break
        if len(values) > MAX_COLUMNS:
            truncated = True
        cells = [_cell(v) for v in values[:MAX_COLUMNS]]
        width = max(width, len(cells))
        rows.append({"number": number, "cells": cells})
    return {"kind": "sheets", "sheets": [{"name": None, "columns": width, "rows": rows}] if rows else [],
            "truncated": truncated}


def _lines(text: str) -> dict[str, Any]:
    lines = [line for line in text.splitlines() if line.strip()]
    return {"kind": "lines", "lines": lines[:MAX_LINES], "truncated": len(lines) > MAX_LINES}


def document_content(path: Path, mime_type: str | None = None) -> dict[str, Any]:
    """The file's content as sheets (Excel, CSV) or lines (Word, text)."""

    kind = ParserService.parser_kind(path, mime_type)
    if kind == "excel":
        return _workbook(path)
    if kind == "csv":
        return _csv(path)
    if kind == "docx":
        try:
            return _lines(docx_text(path))
        except DocxError as exc:
            raise UnreadableContent(str(exc)) from exc
    if kind == "text":
        return _lines(_decoded(path))
    raise UnsupportedContent("This file is shown as pages, not as content.")
