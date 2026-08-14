"""Methodology workbook loader for Subsystem 2 PR1.

Loads the five authoritative S2 methodology workbooks into typed records. This
module intentionally does not execute rules, evaluate conditions, or map S1
approved evidence. It only turns source spreadsheets into validated runtime
objects that later S2 services can consume.
"""

from __future__ import annotations

from collections.abc import Mapping
from pathlib import Path
from typing import Any

from backend.app.domain.methodology import (
    MethodologyBundle,
    MethodologyEdge,
    MethodologySchemaRow,
    MethodologySourceFile,
    VerificationRule,
)

CORE_METHODOLOGY_FILES: Mapping[str, tuple[str, str]] = {
    "gen_schema": ("S2_General_Methodology_Schema_v.1.xlsx", "1_Schema"),
    "scope1_schema": ("S2_Scope1_Data_Schema_v.1.xlsx", "1_Schema"),
    "scope2_schema": ("S2_Scope2_Data_Schema_v.1.xlsx", "S2_Schema"),
    "edges": ("S2_Cross_Layer_Edge_Table_v1.xlsx", "Edge_Table"),
    "rules": ("S2_Verification_Rules_v.1.xlsx", "Verification_Rules"),
}

EXPECTED_ROW_COUNTS: Mapping[str, int] = {
    "gen_schema": 51,
    "scope1_schema": 139,
    "scope2_schema": 89,
    "edges": 114,
    "rules": 89,
}

SCHEMA_REQUIRED_COLUMNS = (
    "Field_ID",
    "Block",
    "Grain",
    "Data_Schema_Field(s)",
    "Calculation_Role",
    "Value_Origin",
    "Requirement_Level",
    "Condition",
    "Vocabulary_Reference",
    "Derived_From",
    "Node_ID",
)

EDGE_REQUIRED_COLUMNS = (
    "Edge_ID",
    "Layer",
    "Scope_Field_ID",
    "Direction",
    "Reference_Type",
    "GEN_Field_ID",
    "Status",
)

RULE_REQUIRED_COLUMNS = (
    "Rule_ID",
    "Layer",
    "Type",
    "Assertion",
    "Applies_To",
    "Rule_Expression",
    "Status",
)


class MethodologyLoaderError(RuntimeError):
    """Base error for methodology loading failures."""


class MethodologySourceNotFoundError(MethodologyLoaderError):
    """Raised when an expected methodology workbook is missing."""


class MethodologySheetError(MethodologyLoaderError):
    """Raised when an expected methodology worksheet is missing."""


class MethodologyColumnError(MethodologyLoaderError):
    """Raised when required worksheet columns are missing."""


class MethodologyDataError(MethodologyLoaderError):
    """Raised when workbook rows fail basic required-value checks."""


def find_default_methodology_dir(repo_root: Path | None = None) -> Path:
    """Return the default methodology directory under reference-data."""

    root = repo_root.resolve() if repo_root is not None else Path(__file__).resolve().parents[3]
    directory = root / "reference-data" / "methodology"
    if not directory.exists() or not directory.is_dir():
        raise MethodologySourceNotFoundError(
            f"Methodology directory not found at: {directory}"
        )
    return directory


def find_default_methodology_sources(
    repo_root: Path | None = None,
) -> tuple[MethodologySourceFile, ...]:
    """Resolve the five authoritative core methodology workbooks."""

    directory = find_default_methodology_dir(repo_root=repo_root)
    sources: list[MethodologySourceFile] = []
    for role, (filename, sheet_name) in CORE_METHODOLOGY_FILES.items():
        path = directory / filename
        if not path.exists() or not path.is_file():
            raise MethodologySourceNotFoundError(
                f"Methodology source for role '{role}' not found at: {path}"
            )
        sources.append(MethodologySourceFile(role=role, path=path, sheet_name=sheet_name))
    return tuple(sources)


def load_default_methodology_bundle(repo_root: Path | None = None) -> MethodologyBundle:
    """Load the default S2 methodology bundle from reference-data."""

    return load_methodology_bundle(find_default_methodology_sources(repo_root=repo_root))


def load_methodology_bundle(
    sources: tuple[MethodologySourceFile, ...],
) -> MethodologyBundle:
    """Load configured methodology sources into one typed bundle."""

    by_role = {source.role: source for source in sources}
    missing_roles = [role for role in CORE_METHODOLOGY_FILES if role not in by_role]
    if missing_roles:
        raise MethodologySourceNotFoundError(
            f"Missing methodology source role(s): {missing_roles}"
        )

    schema_rows: list[MethodologySchemaRow] = []
    schema_rows.extend(_load_schema_rows(by_role["gen_schema"], layer="GEN"))
    schema_rows.extend(_load_schema_rows(by_role["scope1_schema"], layer="S1"))
    schema_rows.extend(_load_schema_rows(by_role["scope2_schema"], layer="S2"))
    edges = _load_edges(by_role["edges"])
    rules = _load_rules(by_role["rules"])

    _validate_expected_count("gen_schema", len([row for row in schema_rows if row.layer == "GEN"]))
    _validate_expected_count("scope1_schema", len([row for row in schema_rows if row.layer == "S1"]))
    _validate_expected_count("scope2_schema", len([row for row in schema_rows if row.layer == "S2"]))
    _validate_expected_count("edges", len(edges))
    _validate_expected_count("rules", len(rules))

    return MethodologyBundle(
        schema_rows=tuple(schema_rows),
        edges=tuple(edges),
        rules=tuple(rules),
        source_files=sources,
    )


def _load_workbook(source: MethodologySourceFile) -> Any:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:  # pragma: no cover - environment dependency
        raise MethodologyLoaderError(
            "openpyxl is required to load methodology workbooks."
        ) from exc

    if not source.path.exists() or not source.path.is_file():
        raise MethodologySourceNotFoundError(
            f"Methodology workbook not found at: {source.path}"
        )
    workbook = load_workbook(source.path, read_only=True, data_only=True)
    if source.sheet_name not in workbook.sheetnames:
        raise MethodologySheetError(
            f"Workbook {source.path.name} missing expected sheet "
            f"'{source.sheet_name}'. Found sheets: {workbook.sheetnames}"
        )
    return workbook


def _load_schema_rows(
    source: MethodologySourceFile,
    *,
    layer: str,
) -> tuple[MethodologySchemaRow, ...]:
    rows = _read_rows(source, required_columns=SCHEMA_REQUIRED_COLUMNS)
    parsed: list[MethodologySchemaRow] = []
    for row_number, row in rows:
        field_id = _required_text(row, "Field_ID", source, row_number)
        scope_value = row.get("Scope_Applicability")
        if scope_value is None:
            scope_value = row.get("Scope")
        parsed.append(
            MethodologySchemaRow(
                field_id=field_id,
                layer=layer,  # type: ignore[arg-type]
                scope_applicability=_optional_text(scope_value),
                block=_optional_text(row.get("Block")),
                grain=_optional_text(row.get("Grain")),
                source_category=_optional_text(row.get("Source_Category")),
                calculation_method=_optional_text(row.get("Calculation_Method")),
                emission_source=_optional_text(row.get("Emission_Source")),
                gas_applicability=_optional_text(row.get("Gas_Applicability")),
                data_schema_fields=_split_multi(row.get("Data_Schema_Field(s)")),
                calculation_role=_optional_text(row.get("Calculation_Role")),
                value_origin=_optional_text(row.get("Value_Origin")),
                requirement_level=_optional_text(row.get("Requirement_Level")),
                condition=_optional_text(row.get("Condition")),
                anchor_type=_optional_text(row.get("Anchor_Type")),
                definition_purpose=_optional_text(row.get("Definition_Purpose")),
                vocabulary_references=_split_multi(row.get("Vocabulary_Reference")),
                source_citation=_optional_text(row.get("Source_Citation")),
                derived_from=_split_multi(row.get("Derived_From")),
                node_id=_optional_text(row.get("Node_ID")),
                notes_flags=_optional_text(row.get("Notes_Flags")),
                source_workbook=source.path.name,
                source_sheet=source.sheet_name,
                source_row_number=row_number,
            )
        )
    return tuple(parsed)


def _load_edges(source: MethodologySourceFile) -> tuple[MethodologyEdge, ...]:
    rows = _read_rows(source, required_columns=EDGE_REQUIRED_COLUMNS)
    parsed: list[MethodologyEdge] = []
    for row_number, row in rows:
        parsed.append(
            MethodologyEdge(
                edge_id=_required_text(row, "Edge_ID", source, row_number),
                layer=_required_text(row, "Layer", source, row_number),
                scope_field_id=_optional_text(row.get("Scope_Field_ID")),
                scope_fields=_split_multi(row.get("Scope_Field(s)")),
                direction=_optional_text(row.get("Direction")),
                reference_type=_optional_text(row.get("Reference_Type")),
                edge_type=_optional_text(row.get("Edge_Type")),
                condition=_optional_text(row.get("Condition")),
                transformation=_optional_text(row.get("Transformation")),
                gen_field_id=_optional_text(row.get("GEN_Field_ID")),
                gen_fields=_split_multi(row.get("GEN_Field(s)")),
                cardinality=_optional_text(row.get("Cardinality")),
                gen_version=_optional_text(row.get("GEN_Version")),
                status=_optional_text(row.get("Status")),
                notes=_optional_text(row.get("Notes")),
                source_workbook=source.path.name,
                source_sheet=source.sheet_name,
                source_row_number=row_number,
            )
        )
    return tuple(parsed)


def _load_rules(source: MethodologySourceFile) -> tuple[VerificationRule, ...]:
    rows = _read_rows(source, required_columns=RULE_REQUIRED_COLUMNS)
    parsed: list[VerificationRule] = []
    for row_number, row in rows:
        parsed.append(
            VerificationRule(
                rule_id=_required_text(row, "Rule_ID", source, row_number),
                layer=_required_text(row, "Layer", source, row_number),
                rule_type=_optional_text(row.get("Type")),
                assertion=_optional_text(row.get("Assertion")),
                applies_to=_split_multi(row.get("Applies_To")),
                grain_key=_optional_text(row.get("Grain_Key")),
                rule_expression=_optional_text(row.get("Rule_Expression")),
                source=_optional_text(row.get("Source")),
                status=_optional_text(row.get("Status")),
                notes=_optional_text(row.get("Notes")),
                source_workbook=source.path.name,
                source_sheet=source.sheet_name,
                source_row_number=row_number,
            )
        )
    return tuple(parsed)


def _read_rows(
    source: MethodologySourceFile,
    *,
    required_columns: tuple[str, ...],
) -> tuple[tuple[int, dict[str, Any]], ...]:
    workbook = _load_workbook(source)
    worksheet = workbook[source.sheet_name]
    raw_headers = next(worksheet.iter_rows(min_row=1, max_row=1, values_only=True))
    headers = [_optional_text(value) for value in raw_headers]
    present = {header for header in headers if header}
    missing = [column for column in required_columns if column not in present]
    if missing:
        raise MethodologyColumnError(
            f"{source.path.name}/{source.sheet_name} missing required column(s): {missing}"
        )

    rows: list[tuple[int, dict[str, Any]]] = []
    for row_number, values in enumerate(
        worksheet.iter_rows(min_row=2, values_only=True),
        start=2,
    ):
        if not any(_optional_text(value) for value in values):
            continue
        row = {
            str(header): value
            for header, value in zip(headers, values)
            if header is not None
        }
        rows.append((row_number, row))
    return tuple(rows)


def _validate_expected_count(role: str, actual: int) -> None:
    expected = EXPECTED_ROW_COUNTS[role]
    if actual != expected:
        raise MethodologyDataError(
            f"Methodology role '{role}' expected {expected} data rows; found {actual}."
        )


def _required_text(
    row: Mapping[str, Any],
    column: str,
    source: MethodologySourceFile,
    row_number: int,
) -> str:
    value = _optional_text(row.get(column))
    if value is None:
        raise MethodologyDataError(
            f"{source.path.name}/{source.sheet_name} row {row_number}: "
            f"required column '{column}' is blank."
        )
    return value


def _optional_text(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    text = str(value).strip()
    return text or None


def _split_multi(value: Any) -> tuple[str, ...]:
    text = _optional_text(value)
    if text is None:
        return ()
    if text in {"-", "N/A", "n/a"}:
        return ()
    parts = [part.strip() for part in text.split(";")]
    return tuple(part for part in parts if part and part not in {"-", "N/A", "n/a"})
