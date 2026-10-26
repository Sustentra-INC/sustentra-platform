"""The Excel workpaper's Summary sheet names the reporting period in the Output screen's words."""

from __future__ import annotations

import io
from datetime import date
from typing import Any

import pytest
from openpyxl import load_workbook

from backend.app.services.workpaper_export_service import build_workbook, format_reporting_period


def _summary(engagement: dict[str, Any]) -> dict[str, Any]:
    content = build_workbook(
        engagement={"name": "FY2026 GHG verification", "client_name": "Acme Foods", "settings": {}, **engagement},
        recalculation={},
        register={},
        items=[],
        exported_by="member@acme.test",
    )
    sheet = load_workbook(io.BytesIO(content))["Summary"]
    return {row[0]: row[1] for row in sheet.iter_rows(min_row=2, values_only=True)}


@pytest.mark.parametrize(
    ("start", "end", "expected"),
    [
        # As the database returns them: dates, or None while Setup has not set one.
        (date(2026, 1, 1), date(2026, 12, 31), "Jan 1, 2026 – Dec 31, 2026"),
        (date(2026, 10, 8), None, "Oct 8, 2026 – end date not set"),
        (None, date(2026, 12, 31), "start date not set – Dec 31, 2026"),
        (None, None, "Reporting period not set"),
        # As ISO text.
        ("2026-10-08", "2026-12-31", "Oct 8, 2026 – Dec 31, 2026"),
    ],
)
def test_the_summary_sheet_names_the_reporting_period_like_the_output_screen(
    start: date | str | None, end: date | str | None, expected: str
) -> None:
    period = _summary({"reporting_period_start": start, "reporting_period_end": end})["Reporting period"]
    assert period == expected
    assert "?" not in period


def test_a_date_it_cannot_read_is_shown_as_it_was_given() -> None:
    assert format_reporting_period("2026-13-40", "soon") == "2026-13-40 – soon"
    assert format_reporting_period("", "") == "Reporting period not set"
