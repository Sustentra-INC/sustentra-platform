from __future__ import annotations

import pytest
from fastapi import HTTPException
from fastapi.testclient import TestClient

from backend.app.api import recalculation as recalculation_api
from backend.app.api.engagements import EngagementUpdate
from backend.app.main import app
from backend.tests.api.conftest import ORG_A, ORG_B
from backend.tests.services.test_engagement_recalculation_service import (
    FakeStore,
    _elec,
    _service,
)

ENGAGEMENT = {"id": "ENG-1", "org_id": ORG_A, "name": "FY2025 GHG", "reporting_period_start": "2025-01-01", "reporting_period_end": "2025-12-31",
              "settings": {"recalculation": {"grid_regions": {"Facility A": "NYUP"}}}}


@pytest.fixture
def wired(monkeypatch: pytest.MonkeyPatch) -> FakeStore:
    store = FakeStore()
    store.add_document("BILL-A", "CT-S2-ELEC", [_elec(1200, "Facility A", "Albany, NY")], org_id=str(ORG_A))

    async def visible(db, engagement_id, user):  # noqa: ANN001, ANN202 - stands in for the Postgres lookup
        if engagement_id != "ENG-1" or user.org_id != ORG_A:
            raise HTTPException(status_code=404, detail="Engagement not found.")
        return ENGAGEMENT

    async def no_db():  # noqa: ANN202
        yield None

    from backend.app.core.db import get_db_session

    monkeypatch.setattr(recalculation_api, "_visible", visible)
    monkeypatch.setattr(recalculation_api, "_service", lambda: _service(store))

    async def no_decisions(db, engagement_id, kind=None):  # noqa: ANN001, ANN202
        return []

    monkeypatch.setattr(recalculation_api.workpaper_item_service, "list_items", no_decisions)

    async def no_audit(*args, **kwargs):  # noqa: ANN002, ANN003, ANN202
        return None

    monkeypatch.setattr(recalculation_api, "write_audit_event", no_audit)
    app.dependency_overrides[get_db_session] = no_db
    yield store
    app.dependency_overrides.pop(get_db_session, None)


def test_recalculation_for_my_engagement(act_as, wired: FakeStore) -> None:
    act_as("org_member", ORG_A)
    with TestClient(app) as client:
        response = client.get("/v1/engagements/ENG-1/recalculation")
    assert response.status_code == 200
    body = response.json()
    location = body["categories"]["purchased_electricity_location_based"]
    assert location["lines"][0]["key"] == "Facility A"
    assert round(location["total_co2e"], 2) == round(1200 * (240.99 + 0.011 * 28 + 0.001 * 265) / 2204.62, 2)
    assert body["gwp_set"] == "AR5" and body["tolerance_pct"] == 5.0
    electricity = body["rollup"]["nodes"]["purchased_electricity_total"]
    assert electricity["co2e_t"] == location["total_co2e"] and electricity["status"] == "complete"


def test_other_org_and_signed_out_cannot_read(act_as, wired: FakeStore) -> None:
    act_as("org_member", ORG_B)
    with TestClient(app) as client:
        assert client.get("/v1/engagements/ENG-1/recalculation").status_code == 404
    act_as(None)
    with TestClient(app) as client:
        assert client.get("/v1/engagements/ENG-1/recalculation").status_code == 401


def test_egrid_zip_lookup(act_as) -> None:
    act_as("org_member", ORG_A)
    with TestClient(app) as client:
        assert client.get("/v1/reference/egrid-zip/12207").json()["subregions"] == ["NYUP"]
        assert client.get("/v1/reference/egrid-zip/10501").json()["subregions"] == ["NYCW", "NYUP"]
        assert client.get("/v1/reference/egrid-zip/99999").status_code == 404
        assert client.get("/v1/reference/egrid-zip/1220").status_code == 422
    act_as(None)
    with TestClient(app) as client:
        assert client.get("/v1/reference/egrid-zip/12207").status_code == 401


def test_reference_lists_the_input_choices(act_as) -> None:
    act_as("org_member", ORG_A)
    with TestClient(app) as client:
        body = client.get("/v1/reference/recalculation").json()
    assert body["gwp_sets"] == ["AR5", "AR6"] and body["default_gwp_set"] == "AR5"
    assert len(body["egrid_subregions"]) == 28 and len(body["stationary_fuels"]) == 63
    assert {"key": "scope1_total", "label": "Scope 1 (Total)", "workbook_row": 22, "unit": "t CO2e"} in \
        body["reported_data_points"]
    assert {"key": "fugitive_total", "label": "Direct emissions from fugitive sources", "workbook_row": 46,
            "unit": "t CO2e", "layer": "scope1", "requirement": "Optional", "check": "recalculated"} in body["data_points"]
    assert {"key": "purchased_steam_total", "label": "Indirect emissions from purchased steam", "workbook_row": 49,
            "unit": "t CO2e", "layer": "scope2", "requirement": "Optional",
            "check": "recalculated"} in body["data_points"]
    assert next(p for p in body["data_points"] if p["key"] == "purchased_cooling_total")["check"] == "not_recalculated"
    assert {"vehicle_type": "Medium- and Heavy-Duty Vehicles", "fuel_type": "Diesel", "on_road": True,
            "by_model_year": True, "first_model_year": 1960, "last_model_year": 2022} in body["vehicle_types"]
    assert {"vehicle_type": "Industrial/Commercial Equipment", "fuel_type": "Diesel", "on_road": False,
            "by_model_year": False, "first_model_year": None, "last_model_year": None} in body["vehicle_types"]
    layout = body["rollup_layout"]
    assert (len(layout["sources"]), len(layout["nodes"]), len(layout["final_outputs"])) == (10, 40, 34)
    assert [step["key"] for step in layout["steps"]] == ["step_1a", "step_1b", "step_2"]  # outputs: a list
    assert any(p["key"] == "purchased_cooling_total" and p["workbook_row"] == 51 for p in body["data_points"])
    assert {"key": "scope1_nf3", "label": "Scope 1 NF3", "workbook_row": 29, "unit": "t", "layer": "scope1",
            "requirement": "Required", "check": "recalculated"} in body["data_points"]


def test_engagement_settings_validate_recalculation_inputs() -> None:
    EngagementUpdate(settings={"recalculation": {"gwp_set": "AR6", "grid_regions": {"Facility A": "NYUP"}}})
    with pytest.raises(ValueError):
        EngagementUpdate(settings={"recalculation": {"gwp_set": "AR4"}})
    EngagementUpdate(settings={"outOfScopeDataPoints": ["process_total"]})
    with pytest.raises(ValueError):
        EngagementUpdate(settings={"outOfScopeDataPoints": ["not_a_data_point"]})


def test_export_is_an_excel_workpaper_of_the_same_numbers(act_as, wired: FakeStore) -> None:
    import io

    from openpyxl import load_workbook

    act_as("org_member", ORG_A)
    with TestClient(app) as client:
        response = client.get("/v1/engagements/ENG-1/export")
    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")
    assert response.headers["content-disposition"].startswith("attachment;")
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Summary", "Final outputs", "Roll-up", "Categories", "Data points", "Working",
                               "Evidence", "Reviewed values", "Requests", "Findings", "Left out"]
    outputs = {row[1]: row for row in book["Final outputs"].iter_rows(min_row=2, values_only=True)}
    assert len(outputs) == 34  # Scope 2 rows 30-37 twice: location- and market-based
    electricity = outputs["Indirect emissions from purchased electricity (location-based)"]
    assert electricity[5] == pytest.approx(1200 * (240.99 + 0.011 * 28 + 0.001 * 265) / 2204.62, rel=1e-3)
    assert outputs["Indirect emissions from purchased cooling"][9] == "Out of scope (not recalculated yet)"
    rollup_lines = [row[1] for row in book["Roll-up"].iter_rows(min_row=2, values_only=True)]
    assert len(rollup_lines) == 40 and "Scope 1 (Total)" in rollup_lines
    categories = {row[0]: row for row in book["Categories"].iter_rows(min_row=2, values_only=True)}
    assert categories["Purchased electricity"][3] == "Calculated" and categories["Process sources"][3] == "Not built yet"
    summary = dict(book["Summary"].iter_rows(min_row=2, values_only=True))
    # Categories not built or without evidence count as 0 (Claire, Q5): the total is complete and
    # the Roll-up sheet says which were counted as 0.
    assert not str(summary["Inventory total, location-based (t CO2e)"]).endswith("(partial)")
    inventory = next(row for row in book["Roll-up"].iter_rows(min_row=2, values_only=True)
                     if str(row[1]).startswith("Inventory total, location-based"))
    assert "Process sources: not built yet" in inventory[8]
    location = next(row for row in book["Data points"].iter_rows(values_only=True)
                    if row[1] == "Scope 2 (Total), location-based")
    assert round(location[4], 2) == round(1200 * (240.99 + 0.011 * 28 + 0.001 * 265) / 2204.62, 2)
    assert any("Table 6" in str(row[4]) for row in book["Working"].iter_rows(values_only=True))
    assert [row[0] for row in book["Evidence"].iter_rows(min_row=2, values_only=True)] == ["BILL-A.pdf"]

    act_as("org_member", ORG_B)
    with TestClient(app) as client:
        assert client.get("/v1/engagements/ENG-1/export").status_code == 404


def test_export_separates_documents_read_from_documents_fully_reviewed(act_as, wired: FakeStore) -> None:
    import io

    from openpyxl import load_workbook

    # BILL-A is accepted; a scanned bill is read but its values still await review.
    wired.add_document("SCAN-1", "CT-S2-ELEC", [_elec(300, "Facility A", "Albany, NY")], org_id=str(ORG_A),
                       decision=None)
    act_as("org_member", ORG_A)
    with TestClient(app) as client:
        response = client.get("/v1/engagements/ENG-1/export")
    assert response.status_code == 200
    summary = dict(load_workbook(io.BytesIO(response.content))["Summary"].iter_rows(min_row=2, values_only=True))
    assert summary["Documents read"] == 2
    assert summary["Documents fully reviewed"] == 1
